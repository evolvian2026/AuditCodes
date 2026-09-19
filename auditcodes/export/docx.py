"""DOCX export built directly from the model with python-docx."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from markdown_it import MarkdownIt

from ..audit.report import AuditReport
from ..exec.languages import LANGUAGES
from ..models import Language, Question
from .html import ExportOptions, _languages

_md = MarkdownIt("commonmark", {"breaks": True, "html": False})
_MONO = "Consolas"
_SEV_COLOR = {"blocker": RGBColor(0xB9, 0x1C, 0x1C), "major": RGBColor(0xB4, 0x53, 0x09), "minor": RGBColor(0x85, 0x4D, 0x0E), "info": RGBColor(0x07, 0x58, 0x85)}


def export_docx(questions: list[Question], reports: dict[str, AuditReport] | None = None, options: ExportOptions | None = None, assets_dir: Path | None = None) -> bytes:
    options = options or ExportOptions()
    reports = reports or {}
    if options.question_ids:
        questions = [q for q in questions if q.id in options.question_ids]
    doc = Document()
    _setup_styles(doc)
    _footer_page_numbers(doc)

    doc.add_paragraph(options.title, style="Title")
    sub = doc.add_paragraph(style="Meta")
    sub.add_run((f"Source: {options.source_name} · " if options.source_name else "") + f"{len(questions)} question{'s' if len(questions) != 1 else ''} · generated {options.generated_at} by AuditCodes")
    headers = ["#", "DBNO", "Title", "Area", "Difficulty", "Tests", "Solutions"] + (["Audit"] if options.include_audit else [])
    rows = []
    for i, q in enumerate(questions, start=1):
        r = reports.get(q.id)
        audit = "not audited"
        if r:
            c = r.counts()
            audit = f"{c['blocker']} blocker, {c['major']} major open; {sum(1 for f in r.findings if f.applied)} applied"
        rows.append([str(q.number or i), q.id, q.title, q.area or "", f"{q.difficulty.value if q.difficulty else ''}{f' ({q.lod})' if q.lod else ''}", f"{len(q.samples)} + {len(q.hidden_tests)}", ", ".join(LANGUAGES[l].display_name for l in Language if l in q.solutions)] + ([audit] if options.include_audit else []))
    _table(doc, headers, rows)

    for i, q in enumerate(questions, start=1):
        doc.add_page_break()
        doc.add_heading(f"Q.{q.number or i} {q.title}", level=1)
        meta = [f"DBNO: {q.id}"]
        if q.area:
            meta.append(f"Area: {q.area}")
        if q.difficulty:
            meta.append(f"Difficulty: {q.difficulty.value.capitalize()}" + (f" (LOD {q.lod})" if q.lod else ""))
        if q.time_limit_seconds:
            meta.append(f"Time limit: {q.time_limit_seconds:g} s")
        doc.add_paragraph("   ".join(meta), style="Meta")

        doc.add_heading("Problem Statement", level=2)
        _markdown(doc, q.description_md, assets_dir)
        if q.input_format_md:
            doc.add_heading("Input Explanation", level=2)
            _markdown(doc, q.input_format_md, assets_dir)
        if q.output_format_md:
            doc.add_heading("Output Explanation", level=2)
            _markdown(doc, q.output_format_md, assets_dir)
        if q.constraints:
            doc.add_heading("Constraints", level=2)
            for c in q.constraints:
                doc.add_paragraph(c.text, style="List Bullet")

        doc.add_heading(f"Sample Test Cases ({len(q.samples)})", level=2)
        for j, t in enumerate(q.samples, start=1):
            doc.add_heading(t.label or f"Example {j}", level=3)
            _io_table(doc, t.stdin, t.stdout)
        if q.sample_explanation_md:
            doc.add_heading("Sample Test Case Explanation", level=2)
            _markdown(doc, q.sample_explanation_md, assets_dir)

        langs = _languages(q, options)
        if options.include_drivers:
            for lang in langs:
                if lang in q.drivers:
                    doc.add_heading(f"Driver Code — {LANGUAGES[lang].display_name}", level=2)
                    _code(doc, q.drivers[lang])
        if options.include_editorial and q.editorial_md:
            doc.add_heading("Editorial", level=2)
            _markdown(doc, q.editorial_md, assets_dir)
        if options.include_solutions:
            for lang in langs:
                if lang in q.solutions:
                    doc.add_heading(f"Code Editorial — {LANGUAGES[lang].display_name}", level=2)
                    _code(doc, q.solutions[lang])
        if options.include_hidden:
            doc.add_heading(f"Hidden Test Cases ({len(q.hidden_tests)})", level=2)
            for j, t in enumerate(q.hidden_tests, start=1):
                bits = [t.label or f"Test Case {j}"]
                if t.difficulty:
                    bits.append(f"Level: {t.difficulty.value.capitalize()}")
                if t.points is not None:
                    bits.append(f"Points: {t.points}")
                if t.origin.value == "generated":
                    bits.append(f"generated · {t.category.value.replace('_', ' ')}")
                doc.add_heading(" · ".join(bits), level=3)
                _io_table(doc, t.stdin, t.stdout)

    if options.include_audit and any(q.id in reports for q in questions):
        doc.add_page_break()
        doc.add_heading("Appendix — Audit Report", level=1)
        doc.add_paragraph("Findings marked execution were established by compiling and running code; findings marked model are reviewed proposals. Applied changes are already reflected in the questions above.", style="Meta")
        for i, q in enumerate(questions, start=1):
            r = reports.get(q.id)
            if not r:
                continue
            doc.add_heading(f"Q.{q.number or i} {q.title} (DBNO {q.id})", level=2)
            line = [f"Audited {r.finished_at or r.started_at}" + (f" · model {r.llm_model}" if r.llm_model else "")]
            for lang, v in r.verifications.items():
                line.append(f"{LANGUAGES[Language(lang)].display_name} editorial: " + ("build failed" if not v.build_ok else f"{v.passed}/{v.total} tests") + (" (after normalisation)" if v.normalized else ""))
            if r.generation and r.generation.get("generated"):
                line.append(f"generated {r.generation['generated']} hidden test(s)")
            doc.add_paragraph(". ".join(line) + ".", style="Meta")
            if not r.findings:
                doc.add_paragraph("No findings.")
                continue
            table = _table(doc, ["Severity", "Rule", "Source", "Component", "Finding", "Status"], [])
            for f in r.sorted_findings:
                cells = table.add_row().cells
                run = cells[0].paragraphs[0].add_run(f.severity.value)
                run.font.color.rgb = _SEV_COLOR[f.severity.value]
                run.bold = f.severity.value in ("blocker", "major")
                cells[1].text = f.rule_id
                cells[2].text = "execution" if f.source.value == "dynamic" else "model"
                cells[3].text = f.component
                p = cells[4].paragraphs[0]
                p.add_run(f.title).bold = True
                p.add_run("\n" + f.message)
                cells[5].text = f.status.value + (" (applied)" if f.applied else "")
            _small_table(table)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- building blocks -------------------------------------------------------------------------


def _setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    if "Code" not in [s.name for s in styles]:
        code = styles.add_style("Code", 1)  # WD_STYLE_TYPE.PARAGRAPH
        code.base_style = normal
        code.font.name = _MONO
        code.font.size = Pt(8.5)
        code.element.rPr.rFonts.set(qn("w:eastAsia"), _MONO)
        pf = code.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
    if "Meta" not in [s.name for s in styles]:
        meta = styles.add_style("Meta", 1)
        meta.base_style = normal
        meta.font.size = Pt(9)
        meta.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _footer_page_numbers(doc: Document) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(8)


def _shade(cell_or_par, fill: str) -> None:
    pr = cell_or_par._element.get_or_add_pPr() if hasattr(cell_or_par, "_element") and cell_or_par._element.tag.endswith("}p") else cell_or_par._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _code(doc: Document, text: str | None) -> None:
    for line in (text or "").rstrip("\n").split("\n"):
        p = doc.add_paragraph(style="Code")
        p.add_run(line if line else " ")
        _shade(p, "F3F4F6")


def _io_table(doc: Document, stdin: str | None, stdout: str | None) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for j, (label, text) in enumerate((("Input", stdin), ("Output", stdout))):
        table.cell(0, j).text = label
        table.cell(0, j).paragraphs[0].runs[0].font.size = Pt(8)
        table.cell(0, j).paragraphs[0].runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        cell = table.cell(1, j)
        cell.paragraphs[0].style = doc.styles["Code"]
        lines = (text or "").rstrip("\n").split("\n")
        cell.paragraphs[0].add_run(lines[0] if lines else "")
        for line in lines[1:]:
            cell.add_paragraph(line, style="Code")
        _shade(cell, "F3F4F6")
    doc.add_paragraph()


def _table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        _shade(cell, "E5E7EB")
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = value
    _small_table(table)
    return table


def _small_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.font.size is None:
                        r.font.size = Pt(9)


def _markdown(doc: Document, text: str | None, assets_dir: Path | None) -> None:
    """Render a Markdown field as Word paragraphs: paragraphs, emphasis, inline code, lists, code fences, images."""
    if not text:
        return
    tokens = _md.parse(text)
    i = 0
    list_style = None
    while i < len(tokens):
        t = tokens[i]
        if t.type == "paragraph_open":
            inline = tokens[i + 1]
            p = doc.add_paragraph(style=list_style or "Normal")
            _inline(p, inline.children or [], doc, assets_dir)
            i += 3
            continue
        if t.type in ("fence", "code_block"):
            _code(doc, t.content)
            i += 1
            continue
        if t.type == "heading_open":
            level = min(int(t.tag[1]), 3)
            p = doc.add_heading("", level=level)
            _inline(p, tokens[i + 1].children or [], doc, assets_dir)
            i += 3
            continue
        if t.type == "bullet_list_open":
            list_style = "List Bullet"
        elif t.type == "ordered_list_open":
            list_style = "List Number"
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            list_style = None
        i += 1


def _inline(p, children, doc: Document, assets_dir: Path | None) -> None:
    bold = italic = code = False
    for c in children:
        if c.type == "text":
            run = p.add_run(c.content)
            run.bold, run.italic = bold, italic
            if code:
                run.font.name = _MONO
        elif c.type == "code_inline":
            run = p.add_run(c.content)
            run.font.name = _MONO
        elif c.type == "strong_open":
            bold = True
        elif c.type == "strong_close":
            bold = False
        elif c.type == "em_open":
            italic = True
        elif c.type == "em_close":
            italic = False
        elif c.type in ("softbreak", "hardbreak"):
            p.add_run().add_break()
        elif c.type == "image":
            src = c.attrGet("src") or ""
            path = (assets_dir / Path(src).name) if assets_dir else None
            if path and path.is_file():
                try:
                    doc.add_picture(str(path), width=Pt(360))
                except Exception:
                    p.add_run(f"[figure: {src}]")
            else:
                p.add_run(f"[figure: {src}]")
