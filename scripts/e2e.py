"""End-to-end check: drive the real server over HTTP and validate every exported artifact.

Unlike the pytest suite (which calls the app in-process), this boots ``auditcodes serve`` as a
subprocess and talks to it over the network, so it also covers packaging: a data file missing from
the wheel fails here and nowhere else. Point ``--python`` at an interpreter that has auditcodes
*installed* (not the source tree) to get that coverage.

    python scripts/e2e.py --python /path/to/venv/bin/python

Model-driven audit stages are skipped when ANTHROPIC_API_KEY is unset; everything else runs.
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent

parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
parser.add_argument("--python", default=sys.executable, help="interpreter that has auditcodes installed")
parser.add_argument("--port", type=int, default=8791)
parser.add_argument("--pdf", default=str(REPO / "tests" / "fixtures" / "coding_ques_sample.pdf"))
parser.add_argument("--workdir", help="where to keep artifacts (default: a temporary directory)")
parser.add_argument("--keep", action="store_true", help="keep the work directory")
args = parser.parse_args()

PYTHON = args.python
BASE = f"http://127.0.0.1:{args.port}"
PDF = Path(args.pdf)
OUT = Path(args.workdir) if args.workdir else Path(tempfile.mkdtemp(prefix="auditcodes-e2e-"))
OUT.mkdir(parents=True, exist_ok=True)
print(f"interpreter: {PYTHON}\nsample:      {PDF}\nartifacts:   {OUT}")

failures = []
step_no = 0


def check(name, cond, detail=""):
    global step_no
    step_no += 1
    status = "PASS" if cond else "FAIL"
    if not cond:
        failures.append(f"{name}: {detail}")
    print(f"  [{status}] {name}" + (f" — {detail}" if detail and not cond else (f" ({detail})" if detail else "")))


def get(path, expect=200):
    req = urllib.request.Request(BASE + path)
    try:
        with urllib.request.urlopen(req, timeout=300) as r:
            return r.status, r.read(), dict(r.headers)
    except urllib.error.HTTPError as e:
        return e.code, e.read(), dict(e.headers)


def post(path, fields=None, files=None, follow=True):
    boundary = "----e2e"
    body = b""
    for k, v in (fields or {}).items():
        body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode()
    for k, (fn, data) in (files or {}).items():
        body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; filename=\"{fn}\"\r\nContent-Type: application/pdf\r\n\r\n".encode() + data + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    req = urllib.request.Request(BASE + path, data=body, headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})

    class NoRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *a, **k):
            return None

    opener = urllib.request.build_opener() if follow else urllib.request.build_opener(NoRedirect)
    try:
        with opener.open(req, timeout=600) as r:
            return r.status, r.read(), dict(r.headers), r.url
    except urllib.error.HTTPError as e:
        return e.code, e.read(), dict(e.headers), e.url


print("\n=== 1. server boots ===")
proc = subprocess.Popen(
    [PYTHON, "-m", "auditcodes", "serve", "--port", str(args.port), "--data-dir", str(OUT / "jobs")],
    stdout=open(OUT / "server.log", "w"), stderr=subprocess.STDOUT,
)
ready = False
for _ in range(60):
    time.sleep(1)
    try:
        if get("/")[0] == 200:
            ready = True
            break
    except Exception:
        pass
check("server responds on /", ready)
if not ready:
    print(Path(OUT / "server.log").read_text()[-3000:])
    proc.kill()
    sys.exit(1)

status, home, _ = get("/")
check("home page lists upload form", b'name="file"' in home and b"Upload a question-bank PDF" in home)
check("static CSS served", get("/static/app.css")[0] == 200)

print("\n=== 2. upload + extraction ===")
st, body, hdrs, url = post("/jobs", files={"file": ("Coding_Ques.pdf", PDF.read_bytes())}, follow=False)
check("upload redirects", st == 303, f"status {st}")
job = hdrs["location"].rstrip("/").rsplit("/", 1)[-1]
check("job id assigned", bool(job), job)
st, page, _ = get(f"/jobs/{job}")
text = page.decode()
check("both questions extracted", "AGENT RA" in text and "Sunehri and her Bag" in text)
check("DBNO ids shown", "105376" in text and "105402" in text)
check("areas extracted", "Greedy Algorithms" in text and "Bit Manipulation" in text)
check("test counts shown", ">10<" in text and ">5<" in text)
check("rejects non-PDF", post("/jobs", files={"file": ("x.txt", b"hello")})[0] == 400)

st, qpage, _ = get(f"/jobs/{job}/q/105376")
qt = qpage.decode()
check("question page renders code", "public static String ancientTranslation" in qt)
check("question page renders markdown", "<strong>s</strong>" in qt)
check("extraction warning surfaced", "typographic characters in code" in qt)
check("hidden tests rendered", "Test Case 10" in qt and "?b?b?b?b?b" in qt)
check("unknown question 404s", get(f"/jobs/{job}/q/999999")[0] == 404)

print("\n=== 3. editing ===")
st, body, _, _ = post(f"/jobs/{job}/q/105376/edit", fields={"path": "title", "value": "AGENT RA (reviewed)"})
check("edit saves", st == 200 and "saved" in body.decode())
check("edit persisted", "AGENT RA (reviewed)" in get(f"/jobs/{job}")[1].decode())
st, body, _, _ = post(f"/jobs/{job}/q/105376/edit", fields={"path": "id", "value": "hacked"})
check("read-only field rejected", "cannot be edited" in body.decode())
st, body, _, _ = post(f"/jobs/{job}/q/105376/edit", fields={"path": "time_limit_seconds", "value": "abc"})
check("bad value rejected", "is not a number" in body.decode())
post(f"/jobs/{job}/q/105376/edit", fields={"path": "title", "value": "AGENT RA"})

print("\n=== 4. audit (execution stages) ===")
st, body, _, _ = post(f"/jobs/{job}/audit")
m = re.search(r'hx-get="([^"]+)"', body.decode())
check("audit-all started", st == 200 and m is not None)
deadline = time.time() + 900
done = False
while time.time() < deadline:
    st, b, _ = get(m.group(1))
    if "hx-trigger" not in b.decode():
        done = True
        break
    time.sleep(2)
check("audit-all completed", done and "Audit finished for 2" in b.decode())

st, qpage, _ = get(f"/jobs/{job}/q/105376")
qt = qpage.decode()
check("Q1 blocker found (SOL-002)", "SOL-002" in qt and "1 blocker" in qt)
check("Q1 duplicate test found (HIDE-001)", "HIDE-001" in qt)
check("patch marked verified by execution", "verified by execution" in qt)
check("editorial verified 13/13", "13/13" in qt)
check("projection offered", "If you accept all" in qt and "Accept all verified patches" in qt)
check("model stages reported as skipped", "no ANTHROPIC_API_KEY" in qt or "model stages skipped" in qt)

st, q2page, _ = get(f"/jobs/{job}/q/105402")
q2t = q2page.decode()
check("Q2 findings (HIDE-001/002/003)", all(r in q2t for r in ("HIDE-001", "HIDE-002", "HIDE-003")))
check("Q2 editorial verified 7/7", "7/7" in q2t)

print("\n=== 5. accepting fixes ===")
st, body, hdrs, _ = post(f"/jobs/{job}/q/105376/findings/accept-verified")
check("accept-verified applied", st == 200 and any(k.lower() == "hx-refresh" for k in hdrs), f"status {st}, headers {sorted(hdrs)}")
st, data, _ = get(f"/jobs/{job}/export.json")
questions = json.loads(data)
q1 = next(q for q in questions if q["id"] == "105376")
check("duplicate hidden test removed", len(q1["hidden_tests"]) == 9, f"{len(q1['hidden_tests'])} tests")
check("curly quotes fixed in editorial", "’" not in q1["solutions"]["java"])
check("DBNO unchanged", q1["id"] == "105376")
qt = get(f"/jobs/{job}/q/105376")[1].decode()
check("findings marked applied", qt.count("accepted") >= 2)
check("projection cleared after accepting", "If you accept all" not in qt)

print("\n=== 6. corrected code actually compiles and passes ===")
st, body, _, _ = post(f"/jobs/{job}/q/105376/verify", fields={"language": "java", "normalize": ""})
m = re.search(r'hx-get="([^"]+)"', body.decode())
deadline = time.time() + 600
while time.time() < deadline:
    st, b, _ = get(m.group(1))
    if "hx-trigger" not in b.decode():
        break
    time.sleep(1)
vt = b.decode()
check("accepted editorial builds without normalisation", "build failed" not in vt, vt[:200])
check("passes all 12 remaining tests", "12/12 passed" in vt, re.search(r"\d+/\d+ passed", vt).group(0) if re.search(r"\d+/\d+ passed", vt) else "?")

print("\n=== 7. exports ===")
for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK"), ("zip", b"PK")):
    st, data, hdrs = get(f"/jobs/{job}/export.{fmt}")
    check(f"{fmt} export downloads", st == 200 and data[:4].startswith(magic[:2]), f"{len(data)} bytes")
    check(f"{fmt} has attachment header", "attachment" in hdrs.get("content-disposition", ""))
    (OUT / f"audited.{fmt}").write_bytes(data)
st, html, _ = get(f"/jobs/{job}/export.html")
(OUT / "audited.html").write_bytes(html)
check("html export renders", st == 200 and b"Appendix" in html)
check("unknown format 404s", get(f"/jobs/{job}/export.xls")[0] == 404)
st, single, _ = get(f"/jobs/{job}/export.html?q=105402")
check("single-question export", b"Sunehri" in single and b"AGENT RA" not in single)
st, nohidden, _ = get(f"/jobs/{job}/export.html?hidden=0&audit=0")
check("export toggles respected", b"Hidden Test Cases" not in nohidden and b"Appendix" not in nohidden)

print("\n=== 8. exported artifacts are correct ===")
import pymupdf
from docx import Document

doc = pymupdf.open(str(OUT / "audited.pdf"))
pdf_text = "\n".join(p.get_text() for p in doc)
check("PDF has pages", len(doc) >= 8, f"{len(doc)} pages")
check("PDF contains both questions", "AGENT RA" in pdf_text and "Sunehri and her Bag" in pdf_text)
check("PDF keeps code indentation", "        StringBuilder s = new StringBuilder(ancientWord);" in pdf_text)
appendix_page = next((i for i, p in enumerate(doc) if "Appendix" in p.get_text()), len(doc))
questions_text = "\n".join(doc[i].get_text() for i in range(appendix_page))
appendix_text = "\n".join(doc[i].get_text() for i in range(appendix_page, len(doc)))
check("PDF reflects accepted fix (no curly quotes in questions)", "’" not in questions_text)
check("appendix still quotes the offending character", "’" in appendix_text)
check("PDF reflects removed duplicate", "Hidden Test Cases (9)" in pdf_text)
check("PDF has audit appendix with decisions", "Appendix" in pdf_text and "SOL-002" in pdf_text and "applied" in pdf_text)
check("PDF has page numbers", re.search(r"\b1 / \d+", pdf_text) is not None)

d = Document(str(OUT / "audited.docx"))
heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
code_lines = [p.text for p in d.paragraphs if p.style.name == "Code"]
check("DOCX has question headings", "Q.1 AGENT RA" in heads and "Q.2 Sunehri and her Bag" in heads)
check("DOCX has appendix", "Appendix — Audit Report" in heads)
check("DOCX keeps code indentation", "        StringBuilder s = new StringBuilder(ancientWord);" in code_lines)
check("DOCX reflects accepted fix", not any("’" in c for c in code_lines))
check("DOCX has tables", len(d.tables) > 15, f"{len(d.tables)} tables")
check("DOCX has page-number field", "PAGE" in d.sections[0].footer._element.xml)

with zipfile.ZipFile(OUT / "audited.zip") as z:
    names = set(z.namelist())
    check("ZIP bundle complete", {"questions.json", "audited.pdf", "audited.docx", "audited.html", "source.pdf", "README.txt"} <= names, str(sorted(names))[:200])
    check("ZIP has audit reports", any(n.startswith("audit/") for n in names))
    zq = json.loads(z.read("questions.json"))
    check("ZIP JSON reflects accepted state", len(next(q for q in zq if q["id"] == "105376")["hidden_tests"]) == 9)
    (OUT / "from-zip.json").write_bytes(z.read("questions.json"))

print("\n=== 9. re-audit of the corrected bank is clean ===")
r = subprocess.run([PYTHON, "-m", "auditcodes", "audit", str(OUT / "from-zip.json"), "-o", str(OUT / "reaudit.json"), "--no-llm"], capture_output=True, text=True, timeout=900)
reports = json.loads((OUT / "reaudit.json").read_text())
r1 = next(x for x in reports if x["question_id"] == "105376")
rules = {f["rule_id"] for f in r1["findings"]}
check("re-audit ran", r.returncode == 0, r.stderr[-300:] if r.returncode else "")
check("SOL-002 gone after fix", "SOL-002" not in rules, str(sorted(rules)))
check("HIDE-001 gone after fix", "HIDE-001" not in rules)
check("editorial now builds unaided", r1["verifications"]["java"]["build_ok"] and not r1["verifications"]["java"]["normalized"])
check("all tests pass", r1["verifications"]["java"]["all_passed"])

print("\n=== 10. CLI paths ===")
r = subprocess.run([PYTHON, "-m", "auditcodes", "extract", str(PDF), "-o", str(OUT / "cli.json"), "--assets", str(OUT / "cli-assets")], capture_output=True, text=True, timeout=300)
check("CLI extract", r.returncode == 0 and len(json.loads((OUT / "cli.json").read_text())) == 2)
r = subprocess.run([PYTHON, "-m", "auditcodes", "export", str(OUT / "cli.json"), "-f", "docx", "-o", str(OUT / "cli.docx")], capture_output=True, text=True, timeout=600)
check("CLI export docx", r.returncode == 0 and (OUT / "cli.docx").stat().st_size > 20000)
r = subprocess.run([PYTHON, "-m", "auditcodes", "export", str(OUT / "jobs" / job), "-f", "pdf", "-o", str(OUT / "cli-job.pdf")], capture_output=True, text=True, timeout=600)
check("CLI export from job dir", r.returncode == 0 and (OUT / "cli-job.pdf").read_bytes()[:4] == b"%PDF")
r = subprocess.run([PYTHON, "-m", "auditcodes", "check-toolchains"], capture_output=True, text=True, timeout=120)
check("CLI check-toolchains", r.returncode == 0 and "MISSING" not in r.stdout)

print("\n=== 11. job lifecycle ===")
st, _, _, _ = post(f"/jobs/{job}/delete", follow=False)
check("job deleted", st == 303)
check("deleted job 404s", get(f"/jobs/{job}")[0] == 404)

proc.terminate()
try:
    proc.wait(timeout=15)
except subprocess.TimeoutExpired:
    proc.kill()
log = (OUT / "server.log").read_text()
check("no tracebacks in server log", "Traceback" not in log, log[-1500:] if "Traceback" in log else "")
check("no 500s in server log", " 500 " not in log, [l for l in log.splitlines() if " 500 " in l][:3])

if not args.keep and not args.workdir:
    shutil.rmtree(OUT, ignore_errors=True)
else:
    print(f"\nartifacts kept in {OUT}")

print("\n" + "=" * 60)
print(f"{step_no - len(failures)}/{step_no} checks passed")
if failures:
    print("\nFAILURES:")
    for f in failures:
        print(" -", f)
    sys.exit(1)
print("END-TO-END: ALL CHECKS PASSED")
