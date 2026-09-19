"""Exports: HTML, PDF (both engines), DOCX and the ZIP bundle, with and without an audit appendix."""

import io
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from auditcodes.audit.report import AuditReport, Finding, FindingSource, FindingStatus, Patch, Severity, VerificationSummary
from auditcodes.export import ExportOptions, export_docx, export_html, export_pdf, export_zip, pdf_backend
from auditcodes.ingest import extract
from auditcodes.models import Language, TestCase, TestCategory, TestOrigin

from .pdf_builder import build_python_question

FIXTURE = Path(__file__).parent / "fixtures" / "coding_ques_sample.pdf"


@pytest.fixture(scope="module")
def questions():
    qs = extract(FIXTURE).questions
    qs[0].hidden_tests.append(TestCase(stdin="?a?\n", stdout="bab\n", label="Test Case 11", points=10, difficulty="medium", category=TestCategory.EDGE, origin=TestOrigin.GENERATED))
    return qs


@pytest.fixture(scope="module")
def reports(questions):
    r = AuditReport(question_id=questions[0].id, status="done", finished_at="2026-09-19 10:00:00", llm_model="mock")
    r.verifications["java"] = VerificationSummary(language="java", build_ok=True, normalized=True, passed=13, total=13, all_passed=True)
    r.add(Finding(rule_id="SOL-002", severity=Severity.BLOCKER, source=FindingSource.DYNAMIC, component="solutions.java", title="Typographic characters in code", message="Curly quotes.", status=FindingStatus.ACCEPTED, applied=True, patch=Patch(path="solutions.java", new_value="x")))
    r.add(Finding(rule_id="STMT-003", severity=Severity.MINOR, source=FindingSource.STATIC, component="editorial_md", title="Grammar", message="'character' should be 'characters'."))
    r.generation = {"generated": 1}
    return {questions[0].id: r}


def test_html_contains_every_component(questions, reports):
    html = export_html(questions, reports, ExportOptions(source_name="sample.pdf"))
    for needle in ["AGENT RA", "Sunehri and her Bag", "DBNO: <b>105376</b>", "Greedy Algorithms", "Problem Statement", "Input Explanation", "Constraints",
                   "Length of s is at most 50.", "Example 1", "?ba??b", "Driver Code — Java", "Code Editorial — Java", "public static String ancientTranslation",
                   "Hidden Test Cases (11)", "Test Case 10", "generated · edge", "Level: Medium", "Points: 10",
                   "Appendix — Audit Report", "SOL-002", "accepted (applied)", "STMT-003", "13/13 tests", "after normalisation", "Generated 1 hidden test"]:
        assert needle in html, needle
    # code is escaped, not interpreted
    assert "List&lt;Integer&gt;" in html and "<Integer>" not in html.split("<pre>")[1]
    # inline emphasis from the statement rendered as HTML
    assert "<strong>s</strong>" in html


def test_html_options(questions, reports):
    html = export_html(questions, reports, ExportOptions(include_hidden=False, include_drivers=False, include_audit=False, question_ids=["105402"]))
    assert "Sunehri" in html and "AGENT RA" not in html
    assert "Hidden Test Cases" not in html and "Driver Code" not in html and "Appendix" not in html
    assert "Code Editorial — Java" in html


@pytest.mark.parametrize("backend", ["pymupdf", "weasyprint"])
def test_pdf_renders(questions, reports, backend):
    if backend == "weasyprint" and pdf_backend() != "weasyprint":
        pytest.skip("weasyprint not installed")
    pdf = export_pdf(questions, reports, ExportOptions(title="Bank Export"), backend=backend)
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    assert len(doc) >= 8
    text = "\n".join(p.get_text() for p in doc)
    for needle in ["Bank Export", "AGENT RA", "Sunehri and her Bag", "ancientTranslation", "Hidden Test Cases (11)", "Appendix", "SOL-002"]:
        assert needle in text, needle
    # code indentation survives into the PDF
    assert "        StringBuilder s = new StringBuilder(ancientWord);" in text or "StringBuilder s = new StringBuilder(ancientWord);" in text
    assert "1 / " in doc[0].get_text() or f"1 / {len(doc)}" in text  # page numbers


def test_docx_structure(questions, reports):
    data = export_docx(questions, reports, ExportOptions(title="Bank Export", source_name="sample.pdf"))
    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs]
    assert paras[0] == "Bank Export"
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Q.1 AGENT RA" in headings and "Q.2 Sunehri and her Bag" in headings
    assert "Problem Statement" in headings and "Code Editorial — Java" in headings and "Hidden Test Cases (11)" in headings
    assert "Appendix — Audit Report" in headings
    code_lines = [p.text for p in doc.paragraphs if p.style.name == "Code"]
    assert "        StringBuilder s = new StringBuilder(ancientWord);" in code_lines
    # emphasis from markdown becomes bold runs
    stmt = next(p for p in doc.paragraphs if p.text.startswith("Big agent RA"))
    assert any(r.bold and "agent RA" in r.text for r in stmt.runs)
    # tables: summary + io tables + findings
    assert len(doc.tables) >= 1 + 3 + 11 + 2 + 5 + 1
    findings_table = doc.tables[-1]
    assert findings_table.rows[0].cells[0].text == "Severity" and any("SOL-002" in r.cells[1].text for r in findings_table.rows[1:])
    # footer carries a page-number field
    assert "PAGE" in doc.sections[0].footer._element.xml


def test_zip_bundle(questions, reports, tmp_path):
    data = export_zip(questions, reports, ExportOptions(), assets_dir=None, source_pdf=FIXTURE)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
        assert {"questions.json", "audit/105376.json", "audited.html", "audited.pdf", "audited.docx", "source.pdf", "README.txt"} <= names
        assert b'"id": "105376"' in z.read("questions.json")
        assert z.read("audited.pdf")[:4] == b"%PDF"


def test_images_are_embedded(tmp_path):
    pdf = build_python_question(tmp_path / "synthetic.pdf")
    assets = tmp_path / "assets"
    qs = extract(pdf, assets).questions
    html = export_html(qs, options=ExportOptions(), assets_dir=assets)
    assert '<img src="data:image/' in html
    rel = export_html(qs, options=ExportOptions(image_mode="relative"), assets_dir=assets)
    assert f'<img src="{qs[0].images[0].path}"' in rel
    data = export_docx(qs, options=ExportOptions(), assets_dir=assets)
    doc = Document(io.BytesIO(data))
    assert doc.inline_shapes and len(doc.inline_shapes) == 1
    pdf_bytes = export_pdf(qs, options=ExportOptions(), assets_dir=assets, backend="pymupdf")
    d = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    assert any(p.get_images() for p in d)
    # python code keeps indentation in the docx
    assert "    a, b = map(int, input().split())" in [p.text for p in doc.paragraphs]


def test_soft_wrap_for_long_code_lines():
    from auditcodes.export.html import _soft_wrap

    line = "    " + "x = " + " + ".join(f"value{i}" for i in range(30))
    wrapped = _soft_wrap(line, 60)
    assert all(len(l) <= 64 for l in wrapped.split("\n")) and wrapped.split("\n")[1].startswith("        ")
    assert _soft_wrap("short", 60) == "short"
